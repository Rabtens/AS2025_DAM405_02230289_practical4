#!/usr/bin/env python3
"""
Build the DAM405 Practical 4 Word report from outputs/results.json.

    python code/build_report.py
    python code/build_report.py --name "Kuenzang" --student-id 0223xxxx

* Every number in the report is read from outputs/results.json (written by
  practical4_automl.py) - nothing is typed in by hand.
* Anything that does not exist yet (e.g. AutoML numbers before you have run
  FLAML) is shown as ⟦yellow highlighted⟧ text and a DRAFT banner is added.
  Re-run this script after the experiment and they all turn into real values.
* Screenshots: save them in screenshots/ with the names listed in README.md
  (fig01_install.png ... fig08_extras.png). They are embedded automatically;
  a missing screenshot appears as a labelled grey placeholder box.
"""
import argparse
import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*\n]*?\*|⟦.+?⟧|`.+?`)")


# --------------------------------------------------------------------------- #
# low-level docx helpers
# --------------------------------------------------------------------------- #
def shade(element_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def cell_shade(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def par_shade(par, fill):
    shade(par._p.get_or_add_pPr(), fill)


def set_cell_borders(cell, color="999999", style="single", sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def add_field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def rich(par, text, size=None, color=None, italic=False, bold=False):
    """Write text into a paragraph. **bold**, `code`, ⟦placeholder⟧ are supported."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.italic = True
        elif part.startswith("⟦") and part.endswith("⟧"):
            r = par.add_run("[" + part[1:-1] + "]")
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x9C, 0x2B, 0x2B)
            continue
        else:
            r = par.add_run(part)
        if bold:
            r.bold = True
        if italic:
            r.italic = True
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return par


class Report:
    def __init__(self, results, screenshots_dir):
        self.R = results
        self.shots = screenshots_dir
        self.fig = 0
        self.tab = 0
        self.missing_shots = []
        self.doc = Document()
        self._setup()

    # ---- document setup ---------------------------------------------------
    def _setup(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        st = d.styles["Normal"]
        st.font.name = "Calibri"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.12
        for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
            h = d.styles[name]
            h.font.name = "Calibri"
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = ACCENT
            h.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
        # footer: module + page number
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("DAM405 · Machine Learning Operations · Practical 4      Page ")
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        pr = fp.add_run()
        pr.font.size = Pt(9)
        add_field(pr, "PAGE")

    # ---- building blocks --------------------------------------------------
    def h1(self, t):
        self.doc.add_heading(t, level=1)

    def h2(self, t):
        self.doc.add_heading(t, level=2)

    def h3(self, t):
        self.doc.add_heading(t, level=3)

    def p(self, text, **kw):
        par = self.doc.add_paragraph()
        rich(par, text, **kw)
        return par

    def bullets(self, items):
        for it in items:
            par = self.doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(2)
            rich(par, it)

    def numbered(self, items):
        for it in items:
            par = self.doc.add_paragraph(style="List Number")
            par.paragraph_format.space_after = Pt(2)
            rich(par, it)

    def code(self, text):
        par = self.doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(0.3)
        par.paragraph_format.space_after = Pt(6)
        par.paragraph_format.line_spacing = 1.0
        par_shade(par, "F2F2F2")
        r = par.add_run(text.strip("\n"))
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(9)

    def output(self, text, label="Output"):
        cap = self.doc.add_paragraph()
        cap.paragraph_format.space_after = Pt(1)
        cap.paragraph_format.keep_with_next = True
        rr = cap.add_run(label)
        rr.italic = True
        rr.font.size = Pt(9)
        rr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        par = self.doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(0.3)
        par.paragraph_format.line_spacing = 1.0
        par_shade(par, "E8F0E8")
        for i, line in enumerate(text.strip("\n").split("\n")):
            if i:
                par.add_run().add_break()
            for seg in re.split(r"(⟦.+?⟧)", line):
                if not seg:
                    continue
                if seg.startswith("⟦"):
                    r = par.add_run("[" + seg[1:-1] + "]")
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    r.bold = True
                else:
                    r = par.add_run(seg)
                r.font.name = "Consolas"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                r.font.size = Pt(9)

    def callout(self, text, fill="FFF4CE", border="E0A800"):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        cell_shade(c, fill)
        set_cell_borders(c, color=border, sz=8)
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        rich(c.paragraphs[0], text, size=10)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def table(self, header, rows, widths_cm, align=None, caption=None):
        if caption:
            self.tab += 1
            cp = self.doc.add_paragraph()
            cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.space_after = Pt(2)
            rich(cp, f"**Table {self.tab}.** {caption}", size=10)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, htxt in enumerate(header):
            c = t.rows[0].cells[i]
            cell_shade(c, "1F3A5F")
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
            rr = c.paragraphs[0].add_run(htxt)
            rr.bold = True
            rr.font.size = Pt(10)
            rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].paragraphs[0].paragraph_format.space_after = Pt(1)
                rich(cells[i].paragraphs[0], str(val), size=10)
                if align and align[i] == "r":
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if not any(header):
            t._tbl.remove(t.rows[0]._tr)
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def screenshot(self, stem, caption, must_show, optional=False):
        """Embed screenshots/<stem>.(png|jpg|jpeg) or draw a placeholder box."""
        self.fig += 1
        found = next((f for ext in ("png", "jpg", "jpeg") for f in [self.shots / f"{stem}.{ext}"] if f.exists()), None)
        if found:
            par = self.doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.keep_with_next = True
            par.add_run().add_picture(str(found), width=Cm(15.5))
        else:
            self.missing_shots.append((stem, must_show, optional))
            t = self.doc.add_table(rows=1, cols=1)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = t.rows[0].cells[0]
            cell_shade(c, "EDEDED")
            set_cell_borders(c, color="888888", style="dashed", sz=8)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraphs[0].paragraph_format.space_before = Pt(18)
            rich(c.paragraphs[0], f"⟦SCREENSHOT PLACEHOLDER{' (optional)' if optional else ''}: save as screenshots/{stem}.png⟧")
            p2 = c.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_after = Pt(18)
            rich(p2, f"Must show: {must_show}", size=9.5, italic=True)
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rich(cap, f"**Figure {self.fig}.** {caption}", size=10)


# --------------------------------------------------------------------------- #
# content
# --------------------------------------------------------------------------- #
def ph(text):
    return f"⟦{text}⟧"


def f4(x):
    return f"{x:.4f}"


def build(rep, args):
    R = rep.R
    env = R.get("environment", {})
    ds = R.get("dataset", {})
    base = R.get("baseline", {})
    sb = R.get("scaled_baseline")
    A = R.get("automl")
    C = R.get("comparison")
    X = R.get("extras")
    S = R.get("saved_model")
    have_auto = A is not None and C is not None
    d = rep.doc

    # ---------------- title block ----------------
    t = d.add_paragraph()
    t.paragraph_format.space_before = Pt(60)
    rich(t, "DAM405 · Machine Learning Operations", size=13, color=RGBColor(0x66, 0x66, 0x66))
    t = d.add_paragraph()
    rich(t, "**Practical 4**", size=30, color=ACCENT)
    t = d.add_paragraph()
    rich(t, "Automated Training & Tuning with an Open-Source AutoML Tool (FLAML)", size=17, color=ACCENT)
    d.add_paragraph()
    rep.table(
        ["", ""],
        [
            ["**Student**", args.name],
            ["**Student ID**", args.student_id if args.student_id else ph("your student ID")],
            ["**Programme**", "B.E. Software Engineering (final year), College of Science and Technology, Royal University of Bhutan"],
            ["**Module / Tutor**", "DAM405 Machine Learning Operations / Kamal Acharya"],
            ["**Date**", args.date],
            ["**Dataset**", ds.get("name", "sklearn built-in breast_cancer")],
        ],
        [4.0, 12.4],
    )
    if not have_auto:
        rep.callout(
            "**DRAFT.** The yellow items in this report are AutoML results that were not available when this file was "
            "generated. Run `python code/practical4_automl.py --extras` on your machine, then run "
            "`python code/build_report.py` again: every yellow item is replaced by the real value and this box disappears."
        )
    d.add_page_break()

    # ---------------- 1. Introduction ----------------
    rep.h1("1. Introduction and objectives")
    rep.p(
        "Building a good machine-learning model normally involves two repetitive decisions: which algorithm to use and "
        "how to set its hyperparameters. **Automated machine learning (AutoML)** automates that search. In this "
        "practical I used **FLAML** (Fast and Lightweight AutoML), an open-source library, to search for a good "
        "classifier inside a fixed time budget, and then checked honestly whether it beats a plain hand-built baseline."
    )
    rep.p("The learning outcomes of the practical, and where each is evidenced in this report:")
    rep.table(
        ["Learning outcome", "Evidence"],
        [
            ["Use an open-source AutoML tool to train and tune a model on a given dataset", "Section 4.3 (Task 3), Figure 4"],
            ["Compare an AutoML result against a simple baseline model", "Sections 4.2 and 4.5 (Tasks 2 and 5), Table 4, Figures 3 and 6"],
            ["Inspect the best model, its hyperparameters and its score", "Section 4.4 (Task 4), Table 3, Figure 5"],
            ["Save the resulting model for later use", "Section 4.6 (Task 6), Figure 7"],
        ],
        [10.2, 6.2],
        caption="Learning outcomes mapped to evidence.",
    )

    # ---------------- 2. Background ----------------
    rep.h1("2. Background")
    rep.h2("2.1 What AutoML automates")
    rep.p(
        "AutoML tools take a dataset, a task type (here classification), a metric (here accuracy) and a compute budget, "
        "and then automatically try different algorithms and hyperparameter settings, evaluate each candidate on "
        "validation data, and return the best one. The human supplies the problem definition and the budget; the tool "
        "supplies the trial-and-error."
    )
    rep.h2("2.2 FLAML")
    rep.p(
        "FLAML is a lightweight open-source AutoML library from Microsoft. Its search strategy is *cost-aware*: it "
        "starts with cheap configurations (small, fast models) and moves towards more expensive ones only when they "
        "look worthwhile, which is why it can find a strong model within a short time budget. In this practical it is "
        "driven through one call, `AutoML.fit(...)`, with `task='classification'`, `metric='accuracy'` and "
        "`time_budget=60` (seconds). The estimators it searches over include gradient-boosted trees (LightGBM, "
        "XGBoost), random forest, extra trees and logistic-regression variants; the exact default list depends on the "
        "FLAML version and the data."
    )
    rep.h2("2.3 Why a baseline is essential")
    rep.p(
        "An AutoML score means little on its own. A simple, well-understood model (here logistic regression) gives a "
        "reference point: if AutoML cannot beat it by a meaningful margin, the extra search cost is not justified."
    )
    rep.h2("2.4 Dataset")
    rep.p(
        f"The example dataset from the practical sheet is scikit-learn's built-in **Breast Cancer Wisconsin "
        f"(Diagnostic)** data: {ds.get('n_samples', 569)} samples, {ds.get('n_features', 30)} numeric features computed "
        f"from digitised images of cell nuclei, and a binary target (0 = malignant, 1 = benign). Class counts: "
        f"{ds.get('class_counts', {'0': 212, '1': 357})}. No dataset was supplied with the practical, so this example "
        f"dataset was used, exactly as the sheet allows."
    )

    # ---------------- 3. Environment ----------------
    rep.h1("3. Environment and tools")
    rep.table(
        ["Item", "Value"],
        [
            ["Operating system", env.get("platform", ph("fill from run"))],
            ["Python", env.get("python", ph("fill from run"))],
            ["scikit-learn", env.get("scikit_learn", ph("fill from run"))],
            ["FLAML", env.get("flaml") or ph("FLAML version - appears after running the AutoML script")],
            ["joblib", env.get("joblib", ph("fill from run"))],
            ["Install command", "`pip install \"flaml[automl]\" scikit-learn`"],
        ],
        [4.5, 11.9],
        caption="Software environment used for the experiment.",
    )
    rep.screenshot("fig01_install", "Installing FLAML and scikit-learn (Task 1, step 1).",
                   "the `pip install \"flaml[automl]\" scikit-learn` command finishing with 'Successfully installed ...'")

    # ---------------- 4. Method and results ----------------
    rep.h1("4. Method, code and results (Tasks 1 to 6)")
    rep.p(
        "All tasks are implemented in `code/practical4_automl.py`, which runs them in order and prints a labelled block "
        "for each one. The snippets below are the sheet's code; the script wraps them with timing, JSON logging and "
        "extra checks. All results are stored in `outputs/results.json`."
    )

    # Task 1
    rep.h2("4.1 Task 1 - Install and load data")
    rep.p("The dataset is loaded and split 80 % / 20 % into training and test sets with a fixed random state so the split is reproducible.")
    rep.code(
        "from sklearn.datasets import load_breast_cancer\n"
        "from sklearn.model_selection import train_test_split\n\n"
        "X, y = load_breast_cancer(return_X_y=True)\n"
        "X_train, X_test, y_train, y_test = train_test_split(\n"
        "    X, y, test_size=0.2, random_state=42)"
    )
    rep.output(
        f"Samples x features: ({ds.get('n_samples', 569)}, {ds.get('n_features', 30)})\n"
        f"Classes / counts : {ds.get('class_counts', {'0': 212, '1': 357})}\n"
        f"X_train / X_test : ({ds.get('n_train', 455)}, {ds.get('n_features', 30)}) ({ds.get('n_test', 114)}, {ds.get('n_features', 30)})"
    )
    rep.p(
        f"The training set has {ds.get('n_train', 455)} rows and the test set {ds.get('n_test', 114)} rows. The test set "
        "is never shown to the AutoML search; it is used once, at the end, for the final comparison."
    )
    rep.screenshot("fig02_data_split", "Task 1 output: dataset shape, class balance and split sizes.",
                   "the 'TASK 1 - Load data and split' block of the script output")

    # Task 2
    rep.h2("4.2 Task 2 - Build a baseline")
    rep.p("Before using AutoML, a plain logistic regression is trained and its test accuracy recorded.")
    rep.code(
        "from sklearn.linear_model import LogisticRegression\n"
        "from sklearn.metrics import accuracy_score\n\n"
        "base = LogisticRegression(max_iter=2000).fit(X_train, y_train)\n"
        "base_acc = accuracy_score(y_test, base.predict(X_test))\n"
        "print('Baseline accuracy:', round(base_acc, 4))"
    )
    rep.output(f"Baseline accuracy: {f4(base['test_accuracy'])}" if base else f"Baseline accuracy: {ph('run the script')}")
    if base:
        rep.p(
            f"The baseline classifies **{base['test_correct']} of {ds.get('n_test', 114)}** test samples correctly "
            f"(accuracy {f4(base['test_accuracy'])}) and trains in about {base['train_seconds']:.2f} s. "
            "Note: scikit-learn prints a `ConvergenceWarning` for this model. The features are on very different scales "
            "and were not standardised (the sheet asks for a *plain* model), so the L-BFGS solver stops at the iteration "
            "limit. The warning is expected, the model is still usable, and it is discussed in Section 5."
        )
    if sb:
        rep.p(
            f"As a supplementary check (not part of the sheet), a baseline with a `StandardScaler` in front reaches "
            f"**{f4(sb['test_accuracy'])}** ({sb['test_correct']} of {ds.get('n_test', 114)}) with no convergence warning. "
            "This is a stricter, fairer human baseline and is used in the discussion."
        )
    rep.screenshot("fig03_baseline", "Task 2 output: baseline accuracy (including the expected ConvergenceWarning).",
                   "the 'TASK 2 - Baseline' block with the line 'Baseline accuracy: ...'")

    # Task 3
    rep.h2("4.3 Task 3 - Run AutoML")
    rep.p(
        "FLAML is given the task, the metric and a time budget of 60 seconds. It searches across estimators and "
        "hyperparameters automatically until the budget is used."
    )
    rep.code(
        "from flaml import AutoML\n\n"
        "automl = AutoML()\n"
        "automl.fit(X_train, y_train, task='classification',\n"
        "           metric='accuracy', time_budget=60)"
    )
    rep.p(
        "In the script the call additionally passes `seed=42` (reproducible search) and `log_file_name` "
        "(writes the full search log to `outputs/flaml_search.log`). Wall-clock time of the search is measured for "
        "the cost analysis."
    )
    if have_auto:
        rep.output(f"AutoML search finished in {A['wall_clock_s']} s (wall clock)")
    else:
        rep.output(f"AutoML search finished in ⟦seconds⟧ s (wall clock)")
    rep.screenshot("fig04_automl_run", "Task 3: FLAML search running and finishing inside the 60 s budget.",
                   "the FLAML log lines (trials / 'Evaluation method') and the final 'AutoML search finished in ... s'")

    # Task 4
    rep.h2("4.4 Task 4 - Inspect the best model")
    rep.code(
        "print('Best estimator :', automl.best_estimator)\n"
        "print('Best config    :', automl.best_config)\n"
        "print('Best CV score  :', 1 - automl.best_loss)"
    )
    if have_auto:
        rep.output(
            f"Best estimator : {A['best_estimator']}\n"
            f"Best config    : {A['best_config']}\n"
            f"Best CV score  : {A['best_cv_score']}"
        )
        rep.table(
            ["Property", "Value"],
            [
                ["Best estimator", f"`{A['best_estimator']}`"],
                ["Best hyperparameters", f"`{A['best_config']}`"],
                ["Best CV score (1 - best_loss)", f4(A["best_cv_score"])],
                ["Best iteration / trial", str(A.get("best_iteration", "n/a"))],
                ["Time to train best config", f"{A.get('best_config_train_time', 'n/a')} s"],
            ],
            [6.0, 10.4],
            caption="The model FLAML selected.",
        )
        rep.p(
            f"FLAML selected **{A['best_estimator']}** with the hyperparameters above. `best_loss` is the search's "
            f"validation loss (for the accuracy metric it is 1 - accuracy), so `1 - best_loss` = "
            f"{f4(A['best_cv_score'])} is the score the search itself measured on its validation folds. This is *not* "
            "the final test score: it is the number the search optimised, so it can be slightly optimistic. For small "
            "datasets like this one FLAML's default resampling normally uses k-fold cross-validation; the 'Evaluation "
            "method' line in the FLAML log (Figure 4) confirms which one was used."
        )
    else:
        rep.output(
            "Best estimator : ⟦e.g. lgbm / xgboost / rf / extra_tree / lrl1 - from your run⟧\n"
            "Best config    : ⟦hyperparameter dictionary from your run⟧\n"
            "Best CV score  : ⟦number from your run⟧"
        )
        rep.table(
            ["Property", "Value"],
            [
                ["Best estimator", ph("from your run")],
                ["Best hyperparameters", ph("from your run")],
                ["Best CV score (1 - best_loss)", ph("from your run")],
                ["Best iteration / trial", ph("from your run")],
                ["Time to train best config", ph("from your run")],
            ],
            [6.0, 10.4],
            caption="The model FLAML selected.",
        )
    rep.screenshot("fig05_best_model", "Task 4 output: best estimator, best hyperparameters and best CV score.",
                   "the 'TASK 4 - Inspect the best model' block, with all three printed lines visible")

    # Task 5
    rep.h2("4.5 Task 5 - Evaluate on the test set")
    rep.code(
        "auto_acc = accuracy_score(y_test, automl.predict(X_test))\n"
        "print('AutoML test accuracy  :', round(auto_acc, 4))\n"
        "print('Baseline test accuracy:', round(base_acc, 4))"
    )
    if have_auto:
        rep.output(
            f"AutoML test accuracy  : {f4(C['automl_test_accuracy'])}\n"
            f"Baseline test accuracy: {f4(C['baseline_test_accuracy'])}\n"
            f"Difference (AutoML - baseline): {C['difference']:+.4f}  = {C['difference_in_samples']:+d} test sample(s) out of {C['n_test']}"
        )
    else:
        rep.output(
            f"AutoML test accuracy  : ⟦number from your run⟧\n"
            f"Baseline test accuracy: {f4(base['test_accuracy'])}\n"
            f"Difference (AutoML - baseline): ⟦from your run⟧"
        )
    rows = [
        ["Baseline (LogisticRegression)", f4(base["test_accuracy"]), f"{base['test_correct']} / {ds.get('n_test', 114)}",
         f"{base['train_seconds']:.2f} s", "-"],
    ]
    if sb:
        rows.append(["Supplementary: scaled LogisticRegression", f4(sb["test_accuracy"]),
                     f"{sb['test_correct']} / {ds.get('n_test', 114)}", "< 1 s", "-"])
    if have_auto:
        rows.append([f"AutoML (FLAML, {A['best_estimator']})", f4(C["automl_test_accuracy"]),
                     f"{C['automl_test_correct']} / {C['n_test']}", f"{C['automl_wall_clock_s']} s", f4(A["best_cv_score"])])
    else:
        rows.append(["AutoML (FLAML)", ph("acc"), ph("n / 114"), ph("s"), ph("cv")])
    rep.table(["Model", "Test accuracy", "Correct", "Training time", "Search CV score"], rows,
              [6.2, 2.6, 2.3, 2.7, 2.6], align=["l", "r", "r", "r", "r"],
              caption="Baseline versus AutoML on the held-out test set.")
    rep.screenshot("fig06_comparison", "Task 5 output: AutoML and baseline test accuracy side by side.",
                   "both accuracy lines, 'AutoML test accuracy' and 'Baseline test accuracy', in the same screenshot")

    # Task 6
    rep.h2("4.6 Task 6 - Save the best model")
    rep.code(
        "import joblib\n"
        "joblib.dump(automl, 'automl_model.joblib')\n"
        "print('Saved automl_model.joblib')"
    )
    if S:
        rep.output(
            f"Saved {S['file']}\nFile size: {S['size_kb']} KB\n"
            f"Reload check - predictions identical after joblib.load(): {S['reload_predictions_identical']}"
        )
        rep.p(
            "To prove the file is usable, the script also reloads it with `joblib.load` and confirms that the reloaded "
            "model produces identical predictions on the test set. To reuse the model later:"
        )
    else:
        rep.output("Saved automl_model.joblib\nFile size: ⟦KB⟧\nReload check - predictions identical after joblib.load(): ⟦True⟧")
        rep.p("The script also reloads the file with `joblib.load` and checks that predictions are identical. To reuse the model later:")
    rep.code(
        "import joblib\n"
        "model = joblib.load('outputs/automl_model.joblib')\n"
        "print(model.predict(X_new))        # X_new must have the same 30 features"
    )
    rep.screenshot("fig07_saved_model", "Task 6: model saved, plus a directory listing showing the file.",
                   "the 'TASK 6' block AND the output of `ls -lh outputs/` (Windows: `dir outputs`) showing automl_model.joblib")

    # Optional experiments
    rep.h2("4.7 Optional experiments - other time budgets and the roc_auc metric")
    rep.p("The sheet suggests re-running with a larger time budget or a different metric. Both were run with `--extras`.")
    if X:
        rows = [[f"{r['time_budget_s']} s", r["best_estimator"], f4(r["best_cv_score"]), f4(r["test_accuracy"]),
                 f"{r['wall_clock_s']} s"] for r in X["budget_sweep"]]
        rep.table(["Time budget", "Best estimator", "Search CV score", "Test accuracy", "Wall clock"], rows,
                  [2.8, 3.6, 3.4, 3.4, 3.2], align=["l", "l", "r", "r", "r"],
                  caption="Effect of the time budget (metric = accuracy).")
        ra = X["roc_auc"]
        rep.table(["Model", "Test ROC-AUC", "Test accuracy"],
                  [["Baseline LogisticRegression", f4(ra["baseline_test_roc_auc"]), f4(base["test_accuracy"])],
                   [f"AutoML, metric = {ra['metric']} ({ra['best_estimator']}, {ra['time_budget_s']} s)",
                    f4(ra["test_roc_auc"]), f4(ra["test_accuracy"])]],
                  [8.4, 4.0, 4.0], align=["l", "r", "r"], caption="Re-running AutoML with roc_auc as the optimisation metric.")
    else:
        rep.table(["Time budget", "Best estimator", "Search CV score", "Test accuracy", "Wall clock"],
                  [[ph("10 s"), ph("-"), ph("-"), ph("-"), ph("-")], [ph("30 s"), ph("-"), ph("-"), ph("-"), ph("-")],
                   [ph("60 s"), ph("-"), ph("-"), ph("-"), ph("-")], [ph("120 s"), ph("-"), ph("-"), ph("-"), ph("-")]],
                  [2.8, 3.6, 3.4, 3.4, 3.2], caption="Effect of the time budget (metric = accuracy).")
        rep.table(["Model", "Test ROC-AUC", "Test accuracy"],
                  [["Baseline LogisticRegression", ph("from run"), f4(base["test_accuracy"])],
                   ["AutoML, metric = roc_auc", ph("from run"), ph("from run")]],
                  [8.4, 4.0, 4.0], caption="Re-running AutoML with roc_auc as the optimisation metric.")
    rep.screenshot("fig08_extras", "Optional experiments output (time-budget sweep and roc_auc run).",
                   "the 'OPTIONAL EXPERIMENTS' block of `python code/practical4_automl.py --extras`", optional=True)

    # ---------------- 5. Discussion ----------------
    rep.h1("5. Discussion and reflection questions")
    n_test = ds.get("n_test", 114)

    rep.h2("Q1. Did AutoML beat your baseline? By how much, and was the extra compute time worth it?")
    if have_auto:
        diff, ns = C["difference"], C["difference_in_samples"]
        ratio = C["automl_wall_clock_s"] / max(C["baseline_train_seconds"], 1e-3)
        if diff > 0:
            verdict = (f"**Yes, marginally.** AutoML reached {f4(C['automl_test_accuracy'])} against the baseline's "
                       f"{f4(C['baseline_test_accuracy'])}, a gain of {diff:+.4f} ({abs(ns)} more of the {n_test} test samples classified correctly).")
        elif diff == 0:
            verdict = (f"**No - it tied.** Both reached {f4(C['automl_test_accuracy'])} ({C['automl_test_correct']} of {n_test} correct).")
        else:
            verdict = (f"**No.** AutoML reached {f4(C['automl_test_accuracy'])} against the baseline's "
                       f"{f4(C['baseline_test_accuracy'])}, a difference of {diff:+.4f} ({abs(ns)} fewer of the {n_test} test samples correct).")
        rep.p(verdict)
        rep.p(
            f"**Cost.** The baseline trained in about {C['baseline_train_seconds']:.2f} s; the AutoML search used "
            f"{C['automl_wall_clock_s']} s of wall-clock time (roughly {ratio:,.0f} times more), plus the extra install and "
            "dependencies. The search CV score was "
            f"{f4(A['best_cv_score'])} versus a test accuracy of {f4(C['automl_test_accuracy'])}; a gap between the two is "
            "normal because the search optimised the former."
        )
    else:
        rep.p(ph("Write after running: did AutoML beat the baseline (0.9561)? by how many test samples? use the numbers from Table 4 - the report generator writes this paragraph automatically once results.json contains the AutoML run."))
    rep.p(
        f"**Interpreting the difference.** The test set has only {n_test} samples, so a single sample is worth "
        f"{100 / n_test:.2f} percentage points. A difference of one to three samples is well inside the noise you would "
        "get from a different random split or a different search run, so it should not be presented as a real "
        "improvement. FLAML's time-budgeted search is also not perfectly repeatable: the number of trials it completes "
        "depends on the speed of the machine, even with a fixed seed."
    )
    if sb:
        rep.p(
            f"**Was the baseline fair?** The sheet's baseline is *unscaled* logistic regression, which triggers a "
            f"convergence warning. Simply standardising the features raises it to {f4(sb['test_accuracy'])} "
            f"({sb['test_correct']} of {n_test}) in a fraction of a second. Any AutoML gain measured against the plain baseline "
            "should therefore be read against this stronger, equally cheap human baseline as well."
        )
    rep.p(
        "**Was the extra compute worth it?** On a small, clean, mostly linearly-separable dataset like this one, a "
        "simple model is already very strong, so the gain from a search is small and the extra minutes of compute buy little. "
        "AutoML earns its cost when the best algorithm is not obvious, the data is harder, or engineer time is the scarce resource. "
        "Here it is best viewed as a check that no obviously better model family was missed."
    )

    rep.h2("Q2. What did AutoML automate, and which decisions still required a human?")
    rep.table(
        ["Automated by FLAML", "Still decided by a human"],
        [
            ["Choosing between algorithm families (boosted trees, forests, logistic regression, ...)", "Defining the problem: is this classification, and what does a wrong prediction cost?"],
            ["Searching each algorithm's hyperparameters", "Choosing the metric (accuracy vs roc_auc vs recall) to match the real-world cost of errors"],
            ["Validating candidates (resampling / cross-validation) and ranking them", "Setting the time budget, i.e. deciding how much compute the answer is worth"],
            ["Stopping the search when the time budget is used", "Splitting data correctly and keeping the test set untouched to avoid leakage"],
            ["Refitting the winner and exposing `predict` / `best_config`", "Data quality, feature engineering, class imbalance and domain knowledge"],
            ["Logging every trial to a file", "Judging whether the result is good enough, fair, explainable and safe to deploy, then saving, versioning and monitoring the model (MLOps)"],
        ],
        [8.2, 8.2],
        caption="Division of work between the tool and the practitioner.",
    )
    rep.p(
        "In short, AutoML automates the *search* but not the *judgement*. Every input it depends on (the data, the "
        "metric, the split, the budget) and every conclusion drawn from its output remained my decision."
    )

    rep.h2("Q3. How does increasing the time budget affect the result and the cost?")
    if X:
        rows = X["budget_sweep"]
        accs = [r["test_accuracy"] for r in rows]
        used_budget = all(r["wall_clock_s"] >= 0.8 * r["time_budget_s"] for r in rows)
        monotone = all(accs[i] <= accs[i + 1] for i in range(len(accs) - 1))
        rep.p(
            f"Table 5 lists runs at budgets of {', '.join(str(r['time_budget_s']) + ' s' for r in rows)}. Test accuracy "
            f"ranged from {min(accs):.4f} to {max(accs):.4f}. "
            + ("Accuracy never got worse as the budget grew, but the gains were small and flattened quickly. "
               if monotone and max(accs) > min(accs)
               else "Accuracy did not change with the budget at all. " if max(accs) == min(accs)
               else "Accuracy did not rise steadily with the budget: extra time produced small, non-monotonic changes. ")
            + ("Wall-clock cost tracked the budget closely, because FLAML keeps searching until the budget is spent."
               if used_budget else "Wall-clock cost was at or below the budget; some runs finished before using all of it.")
        )
    else:
        rep.p(ph("Write after running --extras: compare accuracy at 10 / 30 / 60 / 120 s (Table 5) and note whether it improved, stayed flat or wobbled - the generator writes this from the numbers."))
    rep.p(
        "**General behaviour.** The time budget is a cap on search effort. A longer budget lets FLAML try more "
        "configurations and more expensive models, which usually helps at first and then shows **diminishing returns**: "
        "once a strong configuration has been found, extra time mostly re-discovers the same quality. Cost, however, "
        "grows linearly with the budget (CPU-hours, cloud bill, energy). There is also a subtle risk: the more "
        "candidates that are scored against the same validation data, the more likely the winner is partly a lucky fit "
        "to that data, so the search CV score becomes more optimistic while the test score barely moves. The sensible "
        "approach is to start with a short budget, watch where the improvement curve flattens, and stop there."
    )
    if X:
        ra = X["roc_auc"]
        rep.p(
            f"**Metric.** Optimising `{ra['metric']}` instead of accuracy gave a test ROC-AUC of {f4(ra['test_roc_auc'])} "
            f"(baseline {f4(ra['baseline_test_roc_auc'])}) and test accuracy {f4(ra['test_accuracy'])}, with "
            f"`{ra['best_estimator']}` selected. The metric changes what the search optimises, so it can change the winning model."
        )

    # ---------------- 6. Conclusion ----------------
    rep.h1("6. Conclusion")
    if have_auto:
        rep.p(
            f"A baseline logistic regression reached {f4(base['test_accuracy'])} test accuracy. FLAML, given "
            f"{A['time_budget_s']} s, selected `{A['best_estimator']}` and reached {f4(C['automl_test_accuracy'])} "
            f"({C['difference']:+.4f} vs baseline). The model was saved with joblib and verified by reloading it. The result "
            "shows that AutoML is a fast, low-effort way to check whether a stronger model exists, but that on a small "
            "clean dataset the gain over a sensible simple model may be within noise, and that the human still owns "
            "the metric, the data split and the decision to deploy."
        )
    else:
        rep.p(ph("Conclusion is written automatically from your numbers once the AutoML run is in results.json."))

    # ---------------- Appendices ----------------
    rep.h1("Appendix A - Files and how to reproduce")
    rep.table(
        ["File", "Purpose"],
        [
            ["`code/practical4_automl.py`", "Runs Tasks 1-6 (+ optional experiments) and writes outputs/results.json"],
            ["`code/build_report.py`", "Builds this report from results.json and the screenshots"],
            ["`notebooks/practical4_automl.ipynb`", "Same experiment as a notebook (Jupyter / Colab), one cell per task"],
            ["`outputs/automl_model.joblib`", "The saved AutoML model (Task 6)"],
            ["`outputs/results.json`", "Every number quoted in this report"],
            ["`outputs/flaml_search*.log`", "FLAML's own trial-by-trial search logs"],
            ["`screenshots/`", "Evidence screenshots, fig01 to fig08"],
        ],
        [6.4, 10.0],
        caption="Contents of the submission.",
    )
    rep.code(
        "pip install -r requirements.txt\n"
        "python code/practical4_automl.py --extras\n"
        "python code/build_report.py --name \"Your Name\" --student-id YOUR_ID"
    )
    rep.h1("Appendix B - References")
    rep.bullets(
        [
            "Wang, C., Wu, Q., Weimer, M., & Zhu, E. (2021). FLAML: A Fast and Lightweight AutoML Library. *Proceedings of MLSys 2021*.",
            "FLAML documentation: https://microsoft.github.io/FLAML/",
            "Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. *Journal of Machine Learning Research*, 12, 2825-2830.",
            "Wolberg, W. H., Street, W. N., & Mangasarian, O. L. Breast Cancer Wisconsin (Diagnostic) dataset, UCI Machine Learning Repository.",
        ]
    )


def main():
    ap = argparse.ArgumentParser(description="Build the DAM405 Practical 4 report (.docx)")
    ap.add_argument("--results", default=str(ROOT / "outputs" / "results.json"))
    ap.add_argument("--screenshots", default=str(ROOT / "screenshots"))
    ap.add_argument("--out", default=str(ROOT / "report" / "DAM405_Practical4_Report.docx"))
    ap.add_argument("--name", default="Kuenzang")
    ap.add_argument("--student-id", default="")
    ap.add_argument("--date", default=datetime.now().strftime("%d %B %Y"))
    args = ap.parse_args()

    results = json.loads(Path(args.results).read_text())
    rep = Report(results, Path(args.screenshots))
    build(rep, args)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    rep.doc.save(out)

    have_auto = "automl" in results
    print(f"Report written: {out}")
    print(f"AutoML results present in results.json: {'yes' if have_auto else 'NO (yellow placeholders in report)'}")
    req = [m for m in rep.missing_shots if not m[2]]
    if rep.missing_shots:
        print(f"Screenshots still missing ({len(req)} required, {len(rep.missing_shots) - len(req)} optional):")
        for stem, show, opt in rep.missing_shots:
            print(f"  - screenshots/{stem}.png{' (optional)' if opt else ''}")
    else:
        print("All screenshots embedded.")


if __name__ == "__main__":
    main()
